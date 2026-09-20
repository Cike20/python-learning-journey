from openpyxl import load_workbook
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill

def normalize_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y.%m.%d")

    if isinstance(value, date):
        return value.strftime("%Y.%m.%d")

    return str(value)

def normalize_amount(value):
    if isinstance(value, str):
        value = value.replace(" ", "")
        value = value.replace(",", ".")

    return float(value) 

bank_workbook = load_workbook("bank_transactions.xlsx")
bank_sheet = bank_workbook.active

accounting_workbook = load_workbook("accounting_transactions.xlsx")
accounting_sheet = accounting_workbook.active

bank_transactions = []
accounting_transactions = []
reconciliation_results = []
possible_reconciliation_results = []

for row in bank_sheet.iter_rows(min_row=2, values_only=True):
    date, description, amount = row

    transaction = {
        "date": (date),
        "description": description,
        "amount": (amount)
    }

    bank_transactions.append(transaction)

for row in accounting_sheet.iter_rows(min_row=2, values_only=True):
    date, description, amount = row

    transaction = {
        "date": (date),
        "description": description,
        "amount": (amount)
    }

    accounting_transactions.append(transaction)

for bank_transaction in bank_transactions:
    matched = False
    possible_match = False
    matched_accounting_transaction = None
    
    for accounting_transaction in accounting_transactions:
        if (
            bank_transaction["date"] == accounting_transaction["date"]
            and bank_transaction["description"] == accounting_transaction["description"]
            and bank_transaction["amount"] == accounting_transaction["amount"]
        ):
            matched = True
            matched_accounting_transaction = accounting_transaction
            break

        if (
            bank_transaction["description"] == accounting_transaction["description"]
            and bank_transaction["amount"] == accounting_transaction["amount"]
        ):
            date_difference = abs(
                (bank_transaction["date"] - accounting_transaction["date"]).days
            )
            if date_difference <= 2:
                possible_match = True
                matched_accounting_transaction = accounting_transaction

    if matched:
        status = "Matched"

    elif possible_match:
        status = "Possible Match"

    else:
        status = "Unmatched"

    reconciliation_results.append({
        "bank_date": normalize_date(bank_transaction["date"]),
        "bank_description": bank_transaction["description"],
        "bank_amount": normalize_amount(bank_transaction["amount"]),
        "status": status
    })

    if status == "Possible Match":
        possible_reconciliation_results.append({
            "bank_date": normalize_date(bank_transaction["date"]),
            "bank_description": bank_transaction["description"],
            "bank_amount": normalize_amount(bank_transaction["amount"]),
            "accounting_date": normalize_date(matched_accounting_transaction["date"]),
            "accounting_description": matched_accounting_transaction["description"],
            "accounting_amount": normalize_amount(matched_accounting_transaction["amount"])
        })

huf_format = '#,##0.00 "HUF"'

thin = Side(style="thin")
border = Border(
    left=thin,
    right=thin,
    top=thin,
    bottom=thin
)

workbook = Workbook()

reconciliation_sheet = workbook.active
reconciliation_sheet.title = "Transactions Reconciliation"

possible_match_sheet = workbook.create_sheet("Possible Matches")

reconciliation_sheet.append([
    "Date",
    "Description",
    "Amount",
    "Status"
])

for reconciliation_result in reconciliation_results:
    reconciliation_sheet.append([
        reconciliation_result["bank_date"],
        reconciliation_result["bank_description"],
        reconciliation_result["bank_amount"],
        reconciliation_result["status"]
    ])

possible_match_sheet.append([
    "Bank Date",
    "Bank Description",
    "Bank Amount",
    "Accounting Date",
    "Accounting Description",
    "Accounting Amount"
])

for possible_reconciliation_result in possible_reconciliation_results:
    possible_match_sheet.append([
        possible_reconciliation_result["bank_date"],
        possible_reconciliation_result["bank_description"],
        possible_reconciliation_result["bank_amount"],
        possible_reconciliation_result["accounting_date"],
        possible_reconciliation_result["accounting_description"],
        possible_reconciliation_result["accounting_amount"]
    ])

for sheet in [reconciliation_sheet, possible_match_sheet,]:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

green_fill = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

yellow_fill = PatternFill(
    fill_type="solid",
    fgColor="FFEB9C"
)

red_fill = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)

for row in range(2, reconciliation_sheet.max_row + 1):
    reconciliation_sheet[f"C{row}"].number_format = huf_format
    status_cell = reconciliation_sheet[f"D{row}"]

    if status_cell.value == "Matched":
        status_cell.fill = green_fill

    elif status_cell.value == "Possible Match":
        status_cell.fill = yellow_fill

    elif status_cell.value == "Unmatched":
        status_cell.fill = red_fill

for row in range(2, possible_match_sheet.max_row + 1):
    possible_match_sheet[f"C{row}"].number_format = huf_format
    possible_match_sheet[f"F{row}"].number_format = huf_format

reconciliation_widths = {
    "A": 28,
    "B": 20,
    "C": 16,
    "D": 18,
}

for column, width in reconciliation_widths.items():
    reconciliation_sheet.column_dimensions[column].width = width

possible_match_widths = {
    "A": 16,
    "B": 30,
    "C": 20,
    "D": 18,
    "E": 28,
    "F": 28
}

for column, width in possible_match_widths.items():
    possible_match_sheet.column_dimensions[column].width = width

summary_sheet = workbook.create_sheet("Summary")

matched_count = 0
possible_match_count = 0
unmatched_count = 0

for result in reconciliation_results:
    if result["status"] == "Matched":
        matched_count += 1

    elif result["status"] == "Possible Match":
        possible_match_count += 1

    else:
        unmatched_count += 1

summary_sheet["A1"] = "Bank Reconciliation Summary"

summary_sheet["A3"] = "Status"
summary_sheet["B3"] = "Count"

summary_sheet["A4"] = "Matched"
summary_sheet["B4"] = matched_count

summary_sheet["A5"] = "Possible Match"
summary_sheet["B5"] = possible_match_count

summary_sheet["A6"] = "Unmatched"
summary_sheet["B6"] = unmatched_count

summary_widths = {
    "A": 36,
    "B": 20,
}

for column, width in summary_widths.items():
    summary_sheet.column_dimensions[column].width = width


for cell in summary_sheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

for row in summary_sheet.iter_rows(min_row=3, max_row=summary_sheet.max_row):
    for cell in row:
        cell.border = border

workbook.save("reconciliation_report.xlsx")