from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
import csv
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.formatting.rule import CellIsRule

def normalize_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y.%m.%d")

    if isinstance(value, date):
        return value.strftime("%Y.%m.%d")

    return str(value)

def normalize_amount(value):
    if isinstance(value, str):
        value = value.replace(" ", "")
        value = value.replace(",", ".")

    return float(value)

def process_pdf(file):
    print(f"Processing PDF: {file.name}")
    invoice_data = extract_invoice_data(file)
    all_invoices.append(invoice_data)

def process_excel(file):
    print(f"Processing Excel: {file.name}")
    workbook = load_workbook(file)
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        date, description, category, amount, _ = row

        transaction = {
            "date": normalize_date(date),
            "description": description,
            "category": category,
            "amount": normalize_amount(amount),
            "source_file": file.name
        }

        all_transactions.append(transaction)

def process_csv(file):
    print(f"Processing CSV: {file.name}")
    with open(file, "r") as csv_file:
        reader = csv.reader(csv_file, delimiter=";")

        next(reader)

        for row in reader:
            transaction = {
                "date": normalize_date(row[0]),
                "description": row[1],
                "category": row[2],
                "amount": normalize_amount(row[3]),
                "source_file": file.name
            }

            all_transactions.append(transaction)

def parse_amount(amount_text):
    if amount_text is None:
        return None
    cleaned = amount_text.replace(" ", "")
    cleaned = cleaned.replace(",", ".")
    return float(cleaned)

def extract_invoice_data(pdf_path):
    reader = PdfReader(pdf_path)

    page = reader.pages[0]
    text = page.extract_text()
    lines = text.splitlines()

    invoice_number = None
    invoice_date = None
    amount_due = None
    net_amount = None
    low_tax_value = None
    high_tax_value = None
    gross_amount = None

    for index, line in enumerate(lines):
        if line == "Számla száma":
            invoice_number = lines[index + 1]
        
        if line == "Számla kelte":
            invoice_date = lines[index+1]

        if line.startswith("Számlaérték adó nélkül:"):
            net_amount = line.split(":", 1)[1].strip()

        if line.startswith("5 %-os adóalap"):
            low_tax_value = line.split("Adóérték:", 1)[1].strip()

        if line.startswith("27 %-os adóalap"):
            high_tax_value = line.split("Adóérték:", 1)[1].strip()

        if line.startswith("Számlaérték összesen"):
            gross_amount = line.split("összesen:", 1)[1].strip()

        if line.startswith("Fizetendő:"):
            amount_due = line.split(":", 1)[1].strip()

    net_amount = parse_amount(net_amount)
    gross_amount = parse_amount(gross_amount)
    amount_due = parse_amount(amount_due)
    low_tax_value = parse_amount(low_tax_value)
    high_tax_value = parse_amount(high_tax_value)

    if low_tax_value is None:
        low_tax_value = 0

    if high_tax_value is None:
        high_tax_value = 0

    total_vat = low_tax_value + high_tax_value

    if net_amount is not None and gross_amount is not None:
        difference = abs((net_amount + total_vat) - gross_amount)
        accounting_check = difference < 0.01
    else:
        accounting_check = False

    return {
        "source_file": pdf_path.name,
        "invoice_number": invoice_number,
        "invoice_date": invoice_date,
        "net_amount": net_amount,
        "vat_5": low_tax_value,
        "vat_27": high_tax_value,
        "total_vat": total_vat,
        "gross_amount": gross_amount,
        "amount_due": amount_due,
        "accounting_check": accounting_check
        }

huf_format = '#,##0.00 "HUF"'

thin = Side(style="thin")
border = Border(
    left=thin,
    right=thin,
    top=thin,
    bottom=thin
)

input_folder = Path("input")
all_invoices = []
all_transactions = []

file_counts = {
    "pdf": 0,
    "excel": 0,
    "csv": 0,
    "unsupported": 0
}

for file in input_folder.iterdir():
    if not file.is_file():
        continue

    extension = file.suffix.lower()

    if extension == ".pdf":
        file_counts["pdf"] += 1
        process_pdf(file)

    elif extension == ".xlsx":
        file_counts["excel"] += 1
        process_excel(file)

    elif extension == ".csv":
        file_counts["csv"] += 1
        process_csv(file)

    else:
        file_counts["unsupported"] += 1
        print(f"Unsupported: {file.name}")

workbook = Workbook()

invoice_sheet = workbook.active
invoice_sheet.title = "Invoices"

transaction_sheet = workbook.create_sheet("Transactions")

invoice_sheet.append([
    "Source File",
    "Invoice Number",
    "Invoice Date",
    "Net Amount",
    "VAT 5%",
    "VAT 27%",
    "Total VAT",
    "Gross Amount",
    "Amount Due",
    "Accounting Check"
])

for invoice in all_invoices:
    invoice_sheet.append([
        invoice["source_file"],
        invoice["invoice_number"],
        invoice["invoice_date"],
        invoice["net_amount"],
        invoice["vat_5"],
        invoice["vat_27"],
        invoice["total_vat"],
        invoice["gross_amount"],
        invoice["amount_due"],
        invoice["accounting_check"]
    ])

transaction_sheet.append([
    "Date",
    "Description",
    "Category",
    "Amount",
    "Source File"
])

for transaction in all_transactions:
    transaction_sheet.append([
        transaction["date"],
        transaction["description"],
        transaction["category"],
        transaction["amount"],
        transaction["source_file"]
    ])

for sheet in [invoice_sheet, transaction_sheet]:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

for row in range(2, invoice_sheet.max_row + 1):
    invoice_sheet[f"D{row}"].number_format = huf_format
    invoice_sheet[f"E{row}"].number_format = huf_format
    invoice_sheet[f"F{row}"].number_format = huf_format
    invoice_sheet[f"G{row}"].number_format = huf_format
    invoice_sheet[f"H{row}"].number_format = huf_format
    invoice_sheet[f"I{row}"].number_format = huf_format

for row in range(2, transaction_sheet.max_row + 1):
    transaction_sheet[f"D{row}"].number_format = huf_format

green_fill = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE"
)

red_fill = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)

check_range = f"J2:J{invoice_sheet.max_row}"

invoice_sheet.conditional_formatting.add(
    check_range,
    CellIsRule(
        operator="equal",
        formula=["TRUE"],
        fill=green_fill
    )
)

invoice_sheet.conditional_formatting.add(
    check_range,
    CellIsRule(
        operator="equal",
        formula=["FALSE"],
        fill=red_fill
    )
)

invoice_widths = {
    "A": 28,
    "B": 20,
    "C": 16,
    "D": 18,
    "E": 14,
    "F": 14,
    "G": 16,
    "H": 18,
    "I": 18,
    "J": 18
}

for column, width in invoice_widths.items():
    invoice_sheet.column_dimensions[column].width = width

transaction_widths = {
    "A": 16,
    "B": 30,
    "C": 20,
    "D": 18,
    "E": 28
}

for column, width in transaction_widths.items():
    transaction_sheet.column_dimensions[column].width = width

summary_sheet = workbook.create_sheet("Summary")

total_invoices = len(all_invoices)
total_transactions = len(all_transactions)

invoice_gross_total = 0

for invoice in all_invoices:
    if invoice["gross_amount"] is not None:
        invoice_gross_total += invoice["gross_amount"]

transaction_total = 0

for transaction in all_transactions:
    transaction_total += transaction["amount"]

failed_checks = 0

for invoice in all_invoices:
    if invoice["accounting_check"] == False:
        failed_checks += 1

category_totals = {}

for transaction in all_transactions:
    category = transaction["category"]
    amount = transaction["amount"]

    if category in category_totals:
        category_totals[category] += amount
    else:
        category_totals[category] = amount

summary_sheet["A1"] = "Accounting Intake Summary"

summary_sheet["A3"] = "Metric"
summary_sheet["B3"] = "Value"

summary_sheet["A4"] = "PDF files"
summary_sheet["B4"] = file_counts["pdf"]

summary_sheet["A5"] = "Excel files"
summary_sheet["B5"] = file_counts["excel"]

summary_sheet["A6"] = "CSV files"
summary_sheet["B6"] = file_counts["csv"]

summary_sheet["A7"] = "Unsupported files"
summary_sheet["B7"] = file_counts["unsupported"]

summary_sheet["A9"] = "Invoices extracted"
summary_sheet["B9"] = total_invoices

summary_sheet["A10"] = "Transactions imported"
summary_sheet["B10"] = total_transactions

summary_sheet["A12"] = "Invoice gross total"
summary_sheet["B12"] = invoice_gross_total

summary_sheet["A13"] = "Transaction total"
summary_sheet["B13"] = transaction_total

summary_sheet["B12"].number_format = huf_format
summary_sheet["B13"].number_format = huf_format

summary_sheet["A14"] = "Failed accounting checks"
summary_sheet["B14"] = failed_checks

summary_sheet["A16"] = "Transactions by Category"
summary_sheet["A17"] = "Category"
summary_sheet["B17"] = "Amount"

summary_sheet["A17"].font = Font(bold=True)
summary_sheet["B17"].font = Font(bold=True)

summary_sheet["A1"].font = Font(bold=True, size=14)

summary_sheet["A3"].font = Font(bold=True)
summary_sheet["B3"].font = Font(bold=True)

row_number = 18

for category, amount in category_totals.items():
    summary_sheet[f"A{row_number}"] = category
    summary_sheet[f"B{row_number}"] = amount
    summary_sheet[f"B{row_number}"].number_format = huf_format

    row_number += 1

for row in summary_sheet.iter_rows(
    min_row=16,
    max_row=row_number - 1,
    min_col=1,
    max_col=2
):
    for cell in row:
        cell.border = border

for row in summary_sheet.iter_rows(min_row=3, max_row=13, min_col=1, max_col=2):
    for cell in row:
        cell.border = border

summary_sheet.column_dimensions["A"].width = 28
summary_sheet.column_dimensions["B"].width = 20

workbook.save("accounting_intake_report.xlsx")

print("\n--- File Summary ---")

for file_type, count in file_counts.items():
    print(f"{file_type}: {count}")

print("\n--- PDF Summary ---")

for invoice in all_invoices:
    print(invoice)

print("\n--- Transaction Summary ---")

for transaction in all_transactions:
    print(transaction)