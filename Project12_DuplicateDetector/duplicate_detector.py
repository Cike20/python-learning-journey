from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill

workbook = load_workbook("transactions.xlsx")
sheet = workbook.active

transactions = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    date, supplier, invoice_number, amount = row

    transaction = {
        "date": date,
        "supplier": supplier,
        "invoice_number": invoice_number,
        "amount": amount
    }

    transactions.append(transaction)

invoice_number_counts = {}

for transaction in transactions:
    invoice_number = transaction["invoice_number"]

    if invoice_number in invoice_number_counts:
        invoice_number_counts[invoice_number] += 1
    else:
        invoice_number_counts[invoice_number] = 1

exact_duplicate_counts = {}

for transaction in transactions:
    key = (
        transaction["date"],
        transaction["supplier"],
        transaction["invoice_number"],
        transaction["amount"]
    )

    if key in exact_duplicate_counts:
        exact_duplicate_counts[key] += 1
    else:
        exact_duplicate_counts[key] = 1

duplicate_results = []

for i, transaction_1 in enumerate(transactions):
    for j, transaction_2 in enumerate(transactions):
        if i >= j:
            continue

        if (
            transaction_1["supplier"] == transaction_2["supplier"]
            and transaction_1["amount"] == transaction_2["amount"]
        ):
            date_difference = abs(
                (transaction_1["date"] - transaction_2["date"]).days
            )

            if date_difference <= 3:
                duplicate_results.append({
                    "type": "Possible duplicate",
                    "invoice_1": transaction_1["invoice_number"],
                    "invoice_2": transaction_2["invoice_number"],
                    "details": f"Same supplier and amount, {date_difference} day(s) apart"
                })

for invoice_number, count in invoice_number_counts.items():
    if count > 1:
        duplicate_results.append({
            "type": "Repeated invoice number",
            "invoice_1": invoice_number,
            "invoice_2": invoice_number,
            "details": f"Appears {count} times"
        })

for key, count in exact_duplicate_counts.items():
    if count > 1:
        duplicate_results.append({
            "type": "Exact duplicate",
            "invoice_1": key[2],
            "invoice_2": key[2],
            "details": f"Appears {count} times"
        })

report_workbook = Workbook()
report_sheet = report_workbook.active
report_sheet.title = "Duplicate Review"

report_sheet.append([
    "Type",
    "Invoice 1",
    "Invoice 2",
    "Details"
])

for result in duplicate_results:
    report_sheet.append([
        result["type"],
        result["invoice_1"],
        result["invoice_2"],
        result["details"]
    ])

thin = Side(style="thin")

border = Border(
    left=thin,
    right=thin,
    top=thin,
    bottom=thin
)

red_fill = PatternFill(
    fill_type="solid",
    fgColor="FF6666"
)

orange_fill = PatternFill(
    fill_type="solid",
    fgColor="F4B183"
)

yellow_fill = PatternFill(
    fill_type="solid",
    fgColor="FFD966"
)

for row in range(2, report_sheet.max_row + 1):
    type_cell = report_sheet[f"A{row}"]

    if type_cell.value == "Exact duplicate":
        type_cell.fill = red_fill

    elif type_cell.value == "Repeated invoice number":
        type_cell.fill = orange_fill

    elif type_cell.value == "Possible duplicate":
        type_cell.fill = yellow_fill

for cell in report_sheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")
    cell.border = border

for row in report_sheet.iter_rows():
    for cell in row:
        cell.border = border

report_sheet.column_dimensions["A"].width = 24
report_sheet.column_dimensions["B"].width = 18
report_sheet.column_dimensions["C"].width = 18
report_sheet.column_dimensions["D"].width = 45

report_sheet.freeze_panes = "A2"
report_sheet.auto_filter.ref = report_sheet.dimensions

summary_sheet = report_workbook.create_sheet("Summary")

repeated_count = 0
exact_count = 0
possible_count = 0

for result in duplicate_results:
    if result["type"] == "Repeated invoice number":
        repeated_count += 1

    elif result["type"] == "Exact duplicate":
        exact_count += 1

    elif result["type"] == "Possible duplicate":
        possible_count += 1

summary_sheet["A1"] = "Duplicate Detection Summary"

summary_sheet["A3"] = "Type"
summary_sheet["B3"] = "Count"

summary_sheet["A4"] = "Repeated invoice numbers"
summary_sheet["B4"] = repeated_count

summary_sheet["A5"] = "Exact duplicates"
summary_sheet["B5"] = exact_count

summary_sheet["A6"] = "Possible duplicates"
summary_sheet["B6"] = possible_count

summary_sheet.column_dimensions["A"].width = 30
summary_sheet.column_dimensions["B"].width = 18

report_workbook.save("duplicate_review.xlsx")